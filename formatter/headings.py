"""
Heading & Table of Contents Manager Module
Manages heading hierarchy and builds the publication Table of Contents.
"""

from typing import List, Dict, Any
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


class HeadingManager:
    """Manages heading hierarchy and Table of Contents generation."""

    @staticmethod
    def generate_table_of_contents(
        doc: Document,
        headings: List[Dict[str, Any]],
        insert_position_index: int = 1,
    ) -> None:
        """
        Inserts a clean, publication-formatted Table of Contents page.
        Uses detected CHAPTER_HEADING and SUBHEADING entries without modifying author text.
        """
        if not headings:
            return

        # Find target insertion point (typically after author block or before first chapter)
        target_p = doc.paragraphs[min(insert_position_index, len(doc.paragraphs) - 1)]

        # Add TOC Heading
        toc_title = target_p.insert_paragraph_before("Table of Contents")
        toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_title.paragraph_format.space_before = Pt(24.0)
        toc_title.paragraph_format.space_after = Pt(18.0)
        for run in toc_title.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(16.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)

        # Build TOC Entries
        for h in headings:
            title_text = h.get("text", "").strip()
            level = h.get("final_class", "CHAPTER_HEADING")
            page_num = h.get("estimated_page", 1)

            entry_p = target_p.insert_paragraph_before()
            p_fmt = entry_p.paragraph_format
            p_fmt.line_spacing = 1.15
            p_fmt.space_before = Pt(3.0)
            p_fmt.space_after = Pt(3.0)

            is_chapter = level == "CHAPTER_HEADING"

            if is_chapter:
                p_fmt.left_indent = Cm(0.0)
                p_fmt.space_before = Pt(6.0)
            else:
                p_fmt.left_indent = Cm(0.8)

            # Add Title run
            run_title = entry_p.add_run(title_text)
            run_title.font.name = "Times New Roman"
            run_title.font.size = Pt(11.0 if is_chapter else 10.0)
            run_title.font.bold = is_chapter
            run_title.font.color.rgb = RGBColor(0, 0, 0)

            # Dot leader
            leader_dots = " . " * max(2, (45 - len(title_text) // 2))
            run_dots = entry_p.add_run(f" {leader_dots} ")
            run_dots.font.name = "Times New Roman"
            run_dots.font.size = Pt(9.0)
            run_dots.font.color.rgb = RGBColor(120, 120, 120)

            # Page number
            run_page = entry_p.add_run(str(page_num))
            run_page.font.name = "Times New Roman"
            run_page.font.size = Pt(10.0)
            run_page.font.bold = is_chapter
            run_page.font.color.rgb = RGBColor(0, 0, 0)

        # Add page break after TOC
        page_break_p = target_p.insert_paragraph_before()
        page_break_p.paragraph_format.space_after = Pt(12.0)
        page_break_p.add_run().add_break(docx_enum_break_type())


def docx_enum_break_type():
    from docx.enum.text import WD_BREAK
    return WD_BREAK.PAGE
