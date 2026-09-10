"""
Performance Benchmark Module
Tests processing speed, paragraph throughput, and peak memory across 100, 200, and 400+ page manuscripts.
Saves empirical results directly to performance/results.json.
"""

import os
import sys
import time
import json
import psutil
import tempfile
from typing import Dict, Any, List
from docx import Document
from docx.shared import Pt

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from parser.docx_parser import DocxParser
from rules.confidence import HybridClassifier
from formatter.document_formatter import DocumentFormatter


class BenchmarkRunner:
    """Orchestrates memory and latency profiling across varying manuscript scales."""

    @staticmethod
    def get_memory_usage_mb() -> float:
        process = psutil.Process(os.getpid())
        return process.memory_info().rss / (1024 * 1024)

    @classmethod
    def generate_benchmark_manuscript(
        cls,
        target_pages: int,
        output_path: str,
    ) -> Dict[str, Any]:
        """Synthesizes a realistic document scaling to target page volume."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc = Document()

        # Title
        p_title = doc.add_paragraph()
        p_title.add_run(f"Comprehensive Treatise on Distributed Systems and High Performance Computing ({target_pages} Pages)").bold = True

        p_author = doc.add_paragraph()
        p_author.add_run("Prof. Marcus Vance & Dr. Elena Rostova - MIT Computer Science and AI Lab")

        # Number of chapters based on page volume
        num_chapters = max(3, target_pages // 30)
        paras_per_chapter = max(10, (target_pages * 3) // num_chapters)

        table_count = max(2, target_pages // 20)
        figure_count = max(2, target_pages // 15)

        total_paragraphs = 2
        tables_added = 0
        figures_added = 0

        for ch in range(1, num_chapters + 1):
            p_ch = doc.add_paragraph()
            p_ch.add_run(f"Chapter {ch}: Advanced Algorithmic Patterns and Computational Theory").bold = True
            total_paragraphs += 1

            for sec in range(1, 4):
                p_sub = doc.add_paragraph()
                p_sub.add_run(f"{ch}.{sec} Systematic Empirical Analysis of Asymptotic Scaling").bold = True
                total_paragraphs += 1

                for p_idx in range(paras_per_chapter // 3):
                    p_body = doc.add_paragraph()
                    p_body.add_run(
                        "High performance manuscript processing requires constant time feature evaluation and linear memory scaling. "
                        "By streaming paragraph XML elements and avoiding deep object duplications, the formatting engine maintains consistent throughput. "
                        "Offline statistical classifiers verify syntactic structures without network latency or external cloud service bottlenecks."
                    )
                    total_paragraphs += 1

                # Add figure caption
                if figures_added < figure_count:
                    figures_added += 1
                    p_fig = doc.add_paragraph()
                    p_fig.add_run(f"Figure {figures_added}: Empirical runtime latency versus document size curve {figures_added}.").italic = True
                    total_paragraphs += 1

            # Add table
            if tables_added < table_count:
                tables_added += 1
                tbl = doc.add_table(rows=3, cols=3)
                for r in tbl.rows:
                    for c in r.cells:
                        c.text = "Benchmark Datum"

        # References
        p_ref_h = doc.add_paragraph()
        p_ref_h.add_run("References").bold = True
        total_paragraphs += 1

        for r_i in range(1, 15):
            p_ref = doc.add_paragraph()
            p_ref.add_run(f"[{r_i}] Author, A. (202{r_i % 5}). Scientific computing foundations and typography protocols. Journal of Software Architecture, {r_i}(2), 100-120.")
            total_paragraphs += 1

        doc.save(output_path)
        return {
            "path": output_path,
            "target_pages": target_pages,
            "paragraph_count": total_paragraphs,
            "table_count": tables_added,
            "figure_count": figures_added,
        }

    @classmethod
    def benchmark_scale(
        cls,
        target_pages: int,
        temp_dir: str,
    ) -> Dict[str, Any]:
        """Runs complete parse -> classify -> format workflow and records metrics."""
        input_doc_path = os.path.join(temp_dir, f"bench_{target_pages}_in.docx")
        output_doc_path = os.path.join(temp_dir, f"bench_{target_pages}_out.docx")

        gen_meta = cls.generate_benchmark_manuscript(target_pages, input_doc_path)

        # Measure baseline memory
        mem_start = cls.get_memory_usage_mb()
        peak_mem = mem_start

        start_time = time.perf_counter()

        # Step 1: Parse
        parser = DocxParser(input_doc_path).load()
        peak_mem = max(peak_mem, cls.get_memory_usage_mb())

        # Step 2: Classify
        classifier = HybridClassifier()
        classifications = classifier.classify_document(parser.paragraphs)
        peak_mem = max(peak_mem, cls.get_memory_usage_mb())

        # Step 3: Format
        formatter = DocumentFormatter(input_doc_path, output_doc_path, include_toc=True)
        formatter.format(classifications=classifications)
        peak_mem = max(peak_mem, cls.get_memory_usage_mb())

        end_time = time.perf_counter()
        elapsed_sec = round(end_time - start_time, 3)

        # Count detected structures
        detected_counts: Dict[str, int] = {}
        for c in classifications:
            cls_name = c["final_class"]
            detected_counts[cls_name] = detected_counts.get(cls_name, 0) + 1

        total_paras = len(parser.paragraphs)
        paras_per_sec = round(total_paras / max(elapsed_sec, 0.001), 1)

        # Clean up temporary files
        try:
            if os.path.exists(input_doc_path):
                os.remove(input_doc_path)
            if os.path.exists(output_doc_path):
                os.remove(output_doc_path)
        except Exception:
            pass

        return {
            "page_scale": target_pages,
            "paragraphs_processed": total_paras,
            "tables_processed": len(parser.tables),
            "figures_processed": len(parser.images) + detected_counts.get("CAPTION", 0),
            "elapsed_time_seconds": elapsed_sec,
            "throughput_paragraphs_per_sec": paras_per_sec,
            "peak_memory_mb": round(peak_mem, 1),
            "memory_delta_mb": round(peak_mem - mem_start, 1),
            "detected_structures": detected_counts,
        }


def run_benchmark(output_json: str = "performance/results.json") -> Dict[str, Any]:
    """Runs complete 100, 200, and 400+ page benchmark suite."""
    print("=" * 60)
    print("IntelliFormat Performance Benchmark (100, 200, 400+ pages)")
    print("=" * 60)

    scales = [100, 200, 400]
    results_list = []

    with tempfile.TemporaryDirectory() as temp_dir:
        for scale in scales:
            print(f"\n[Running Benchmark: {scale} Pages]...")
            res = BenchmarkRunner.benchmark_scale(scale, temp_dir)
            print(f" -> Processed {res['paragraphs_processed']} paragraphs in {res['elapsed_time_seconds']}s ({res['throughput_paragraphs_per_sec']} paras/sec)")
            print(f" -> Peak Memory: {res['peak_memory_mb']} MB (Delta: +{res['memory_delta_mb']} MB)")
            results_list.append(res)

    suite_summary = {
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "scales_tested": scales,
        "results": results_list,
        "system_info": {
            "python_version": sys.version.split()[0],
            "cpu_count": psutil.cpu_count(logical=True),
            "total_physical_memory_mb": round(psutil.virtual_memory().total / (1024 * 1024), 1),
        },
    }

    os.makedirs(os.path.dirname(output_json), exist_ok=True)
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(suite_summary, f, indent=2)

    print(f"\nAll benchmark results saved to {output_json}")
    return suite_summary


if __name__ == "__main__":
    run_benchmark()
